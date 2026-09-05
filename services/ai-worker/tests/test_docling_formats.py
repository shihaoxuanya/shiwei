from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation

from shiwei_ai.ingestion.parser import DocumentParser


def test_docling_extracts_docx_structure(tmp_path: Path) -> None:
    path = tmp_path / "oracle.docx"
    document = Document()
    document.add_heading("Oracle 恢复记录", level=1)
    document.add_paragraph("恢复以后 PDB 处于 MOUNTED 状态。")
    document.save(path)

    parsed = DocumentParser().parse(path, path.name)

    assert parsed.parser == "docling"
    assert "PDB" in "\n".join(section.plain_text() for section in parsed.sections)
    assert any("Oracle" in " ".join(section.heading_path) for section in parsed.sections)


def test_docling_extracts_pptx_slide_location(tmp_path: Path) -> None:
    path = tmp_path / "restore.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "恢复验证"
    slide.placeholders[1].text = "ALTER PLUGGABLE DATABASE OPEN"
    presentation.save(path)

    parsed = DocumentParser().parse(path, path.name)

    assert "PLUGGABLE" in "\n".join(section.plain_text() for section in parsed.sections)
    assert any(section.slide_number is not None for section in parsed.sections)


def test_docling_extracts_xlsx_cells(tmp_path: Path) -> None:
    path = tmp_path / "checklist.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "恢复检查"
    worksheet.append(["项目", "结果"])
    worksheet.append(["PDB", "OPEN"])
    workbook.save(path)

    parsed = DocumentParser().parse(path, path.name)

    content = "\n".join(section.plain_text() for section in parsed.sections)
    assert "PDB" in content
    assert "OPEN" in content


def test_docling_ocr_extracts_text_from_image(tmp_path: Path) -> None:
    path = tmp_path / "error.png"
    image = Image.new("RGB", (900, 220), "white")
    draw = ImageDraw.Draw(image)
    draw.text((50, 80), "ORACLE ERROR ORA-01034", fill="black", font_size=48)
    image.save(path)

    parsed = DocumentParser().parse(path, path.name)

    content = "\n".join(section.plain_text() for section in parsed.sections)
    assert "ORA-01034" in content.replace(" ", "")
